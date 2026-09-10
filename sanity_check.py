"""
Developer Sanity Check Script for CLAUSE Enterprise RAG (Phase 1)
Validates the end-to-end flow: Document -> Extraction -> Cleaning -> Chunks -> Metadata & Stats.
"""

import sys
import tempfile
from pathlib import Path
from reportlab.pdfgen import canvas
from docx import Document as DocxDocument

from src.document_loader import load_document, load_documents_from_directory
from src.chunking import chunk_documents, build_embedding_text
from src.stats import calculate_document_stats


def run_sanity_check():
    print("=" * 60)
    print("CLAUSE SANITY CHECK - PHASE 1 FOUNDATION & PIPELINE")
    print("=" * 60)

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)

        # 1. Create a sample PDF
        pdf_path = tmp_path / "enterprise_policy.pdf"
        c = canvas.Canvas(str(pdf_path))
        c.drawString(100, 750, "CLAUSE Enterprise Security & Remote Work Policy")
        c.drawString(100, 730, "Section 1: Overview")
        c.drawString(100, 710, "All employees are required to comply with remote access standards.")
        c.showPage()
        c.drawString(100, 750, "Section 2: Data Encryption")
        c.drawString(100, 730, "Data at rest and in transit must be encrypted using AES-256.")
        c.showPage()
        c.save()

        # 2. Create a sample DOCX
        docx_path = tmp_path / "operations_guide.docx"
        doc = DocxDocument()
        doc.add_heading("Standard Operating Procedures", level=1)
        doc.add_paragraph("This document outlines standard operational workflows for all departments.")
        doc.add_heading("Incident Management", level=2)
        doc.add_paragraph("High priority incidents must be escalated within 15 minutes of occurrence.")
        doc.save(str(docx_path))

        # 3. Document Extraction & Cleaning
        print("\n[Step 1] Loading documents from temporary directory...")
        pages = load_documents_from_directory(tmp_path)
        print(f"-> Extracted {len(pages)} total page(s) across PDF and DOCX.")

        for p in pages:
            print(f"   - Doc: '{p.document_name}' | Type: {p.file_type} | Page: {p.page_number} | Length: {len(p.text)} chars")

        # 4. Context-Aware Chunking
        print("\n[Step 2] Chunking extracted pages...")
        chunks = chunk_documents(pages, chunk_size=200, chunk_overlap=40)
        print(f"-> Generated {len(chunks)} chunk(s).")

        for i, chunk in enumerate(chunks[:3]):
            print(f"\n--- Chunk {i} [ID: {chunk.chunk_id}] ---")
            print(f"Document: {chunk.document_name} | Page: {chunk.page_number} | Section: {chunk.section}")
            print(f"Text Snippet: {chunk.text[:80]}...")
            print(f"Embedding Context Sample:\n{build_embedding_text(chunk)[:120]}...")

        # 5. Document Statistics
        print("\n[Step 3] Computing document statistics...")
        stats = calculate_document_stats(pages, chunks)
        print("-> Document Statistics Report:")
        for key, val in stats.items():
            print(f"   - {key}: {val}")

        print("\n" + "=" * 60)
        print("SANITY CHECK SUCCESSFUL: Extraction -> Cleaning -> Chunking -> Metadata pipeline operational.")
        print("=" * 60)


if __name__ == "__main__":
    try:
        run_sanity_check()
    except Exception as e:
        print(f"\n[FAIL] Sanity check failed: {e}", file=sys.stderr)
        sys.exit(1)
