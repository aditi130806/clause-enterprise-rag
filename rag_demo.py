"""
Developer Integration & Grounded RAG Demo for CLAUSE Enterprise RAG.

Evaluates end-to-end question answering across 6 key enterprise test queries using live Gemini LLM service.
"""

import tempfile
import time
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import docx

from src.rag_pipeline import build_and_index_documents, answer_question
from src.llm import GeminiLLMService


def create_synthetic_pdf(filepath: Path, title: str, pages_content: list[tuple[str, str]]):
    """Create a multi-page PDF using reportlab."""
    c = canvas.Canvas(str(filepath), pagesize=letter)
    for idx, (section, text) in enumerate(pages_content, start=1):
        c.setFont("Helvetica-Bold", 14)
        c.drawString(50, 750, f"{title} - Page {idx}")
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, 725, f"Section: {section}")
        c.setFont("Helvetica", 10)

        text_object = c.beginText(50, 700)
        for line in text.split("\n"):
            text_object.textLine(line)
        c.drawText(text_object)
        c.showPage()
    c.save()


def create_synthetic_docx(filepath: Path, title: str, sections: list[tuple[str, str]]):
    """Create a DOCX document using python-docx."""
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for heading, text in sections:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(text)
    doc.save(str(filepath))


def build_synthetic_corpus(doc_dir: Path):
    """Generate 4 realistic synthetic policy documents for evaluation."""
    # 1. Remote Work Policy PDF
    create_synthetic_pdf(
        doc_dir / "remote_work_policy.pdf",
        "CLAUSE Enterprise Remote Work Policy",
        [
            (
                "General Remote Work Guidelines",
                "Policy Identifier: HR-RW-017\n"
                "Section 1: General Remote Work Guidelines\n"
                "Eligible full-time employees may work remotely up to two days per week with prior manager approval. "
                "Core business hours (9 AM to 5 PM EST) must be maintained regardless of physical working location. "
                "Employees are required to maintain a secure home workspace environment.",
            ),
            (
                "Probationary Remote Work Requirements",
                "Section 2: Probationary Remote Work Requirements\n"
                "Probationary employees during their initial 90-day period require both direct manager and HR department approval "
                "before initiating any remote work arrangements. Approval is granted on a case-by-case basis depending on role performance.",
            ),
        ],
    )

    # 2. Annual Leave Policy DOCX
    create_synthetic_docx(
        doc_dir / "annual_leave_policy.docx",
        "CLAUSE Annual Leave & Time Off Policy",
        [
            (
                "Annual Leave Entitlement",
                "All full-time employees receive an annual leave entitlement of 20 paid leave days per calendar year. "
                "Leave requests exceeding 5 consecutive business days require submission at least two weeks in advance. "
                "Unused leave up to 5 days may carry over into the next fiscal year.",
            ),
        ],
    )

    # 3. Expense Reimbursement Policy PDF
    create_synthetic_pdf(
        doc_dir / "expense_reimbursement_policy.pdf",
        "CLAUSE Expense Reimbursement Policy",
        [
            (
                "Business Expense Standards",
                "Section 1: Business Expense Standards\n"
                "Itemized receipt requirements apply to all individual business expense reimbursements exceeding $25.00 USD. "
                "Expense reports must be submitted within 30 calendar days of incurring the expense. "
                "Travel meal expenses are capped at $75 per day.",
            ),
        ],
    )

    # 4. Information Security Policy PDF
    create_synthetic_pdf(
        doc_dir / "information_security_policy.pdf",
        "CLAUSE Corporate Information Security Policy",
        [
            (
                "Authentication & Password Security",
                "Policy Identifier: SEC-AUTH-004\n"
                "Section 1: Authentication & Password Security\n"
                "Multi-factor authentication (MFA) is strictly mandatory for all remote system access, VPN connections, "
                "and enterprise cloud portals. Passwords must be a minimum of 16 characters in length and updated every 90 days. "
                "Hardware security tokens or authenticator apps are required for primary MFA verification.",
            ),
        ],
    )


def run_rag_demo():
    print("=" * 80)
    print("CLAUSE ENTERPRISE RAG - GROUNDED QA BACKEND DEMO")
    print("=" * 80)

    with tempfile.TemporaryDirectory() as tmp_dir:
        doc_dir = Path(tmp_dir) / "docs"
        vstore_dir = Path(tmp_dir) / "vstore"
        doc_dir.mkdir()
        vstore_dir.mkdir()

        print("\n[Step 1] Building synthetic policy corpus...")
        build_synthetic_corpus(doc_dir)

        print("[Step 2] Executing indexing service...")
        retriever, stats = build_and_index_documents(
            documents_dir=doc_dir,
            vectorstore_dir=vstore_dir,
        )

        print("\n--- Indexing Report ---")
        for k, v in stats.items():
            print(f"   - {k}: {v}")

        # Initialize Live Gemini LLM Service directly (no fallback mocks in runtime demo)
        llm_service = GeminiLLMService()

        print(f"\n[Step 3] Initializing Gemini LLM Service...")
        print(f"   - Model: {llm_service.model_name}")

        # Test Queries
        queries = [
            (
                "1. Probationary Remote Work",
                "Can an employee work from home while on probation?",
            ),
            (
                "2. Specific Policy ID Lookup",
                "What does HR-RW-017 say?",
            ),
            (
                "3. Information Security Query",
                "What are the MFA requirements?",
            ),
            (
                "4. Expense Policy Query",
                "What receipts are required for expenses?",
            ),
            (
                "5. Unsupported Query (Hallucination Test)",
                "Does the company provide employees with free company cars?",
            ),
            (
                "6. Multi-Document Security & Remote Work Query",
                "Summarize the MFA security requirements and remote work approval rules.",
            ),
        ]

        print("\n" + "=" * 80)
        print("EXECUTING GROUNDED QA BACKEND EVALUATION")
        print("=" * 80)

        for idx, (label, question) in enumerate(queries):
            if idx > 0:
                time.sleep(6.0)  # Pacing API calls for rate-limit stability

            print(f"\n" + "-" * 80)
            print(f"[{label}]")
            print(f"QUESTION: '{question}'")
            print("-" * 80)

            response = answer_question(
                question=question,
                retriever=retriever,
                llm_service=llm_service,
            )

            print(f"ANSWER:\n{response.answer}\n")
            print(f"STATUS -> Refused: {response.refused} | Grounded: {response.grounded} | Confidence: {response.confidence:.2f}")

            if response.sources:
                print("CITED SOURCES:")
                for src in response.sources:
                    page_str = f"P.{src.page_number}" if src.page_number else "N/A"
                    sec_str = src.section or "N/A"
                    print(f"   - [{src.source_id}] {src.document_name} ({page_str} | {sec_str}) -> Excerpt: \"{src.excerpt[:80]}...\"")
            else:
                print("CITED SOURCES: None")

            if response.metadata.get("refusal_stage"):
                print(f"REFUSAL DIAGNOSTICS: Stage='{response.metadata['refusal_stage']}' | Reason='{response.metadata.get('refusal_reason')}'")

    print("\n" + "=" * 80)
    print("GROUNDED QA DEMO COMPLETED.")
    print("=" * 80)


if __name__ == "__main__":
    run_rag_demo()
