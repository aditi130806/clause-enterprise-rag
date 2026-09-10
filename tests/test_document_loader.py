import tempfile
from pathlib import Path
import pytest
from docx import Document as DocxDocument
from reportlab.pdfgen import canvas

from src.document_loader import (
    load_pdf,
    load_docx,
    load_document,
    load_documents_from_directory,
    DocumentLoadingError,
    UnsupportedFormatError,
    CorruptDocumentError,
)
from src.chunking import chunk_documents
from src.stats import calculate_document_stats


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture
def sample_pdf_path(temp_dir):
    pdf_file = temp_dir / "sample_policy.pdf"
    c = canvas.Canvas(str(pdf_file))
    
    # Page 1
    c.drawString(100, 750, "CLAUSE Enterprise Policy Document")
    c.drawString(100, 730, "This is page one of the enterprise remote work policy.")
    c.showPage()
    
    # Page 2
    c.drawString(100, 750, "Section 2: Compliance Requirements")
    c.drawString(100, 730, "Employees must complete security training annually.")
    c.showPage()
    
    c.save()
    return pdf_file


@pytest.fixture
def sample_docx_path(temp_dir):
    docx_file = temp_dir / "sample_manual.docx"
    doc = DocxDocument()
    doc.add_heading("Employee Operations Manual", level=1)
    doc.add_paragraph("This manual contains operational workflows for department heads.")
    doc.add_heading("Safety Guidelines", level=2)
    doc.add_paragraph("All workplace safety procedures must be adhered to at all times.")
    doc.save(str(docx_file))
    return docx_file


def test_load_pdf_pages(sample_pdf_path):
    pages = load_pdf(sample_pdf_path)
    assert len(pages) == 2
    
    page1 = pages[0]
    assert page1.document_name == "sample_policy.pdf"
    assert page1.file_type == "pdf"
    assert page1.page_number == 1
    assert "enterprise remote work policy" in page1.text
    assert len(page1.document_id) == 16
    assert page1.metadata["total_pages"] == 2

    page2 = pages[1]
    assert page2.page_number == 2
    assert "Compliance Requirements" in page2.text


def test_load_docx(sample_docx_path):
    pages = load_docx(sample_docx_path)
    assert len(pages) == 1
    
    page = pages[0]
    assert page.document_name == "sample_manual.docx"
    assert page.file_type == "docx"
    assert page.page_number is None  # Honest representation for DOCX
    assert "Operations Manual" in page.text
    assert "Safety Guidelines" in page.metadata["headings"]


def test_unsupported_format(temp_dir):
    txt_file = temp_dir / "unsupported.txt"
    txt_file.write_text("Hello world")

    with pytest.raises(UnsupportedFormatError):
        load_document(txt_file)


def test_corrupt_file_handling(temp_dir):
    corrupt_pdf = temp_dir / "corrupt.pdf"
    corrupt_pdf.write_bytes(b"NOT A REAL PDF CONTENT 1234567890")

    with pytest.raises(CorruptDocumentError):
        load_pdf(corrupt_pdf)


def test_missing_file():
    with pytest.raises(DocumentLoadingError):
        load_pdf("non_existent_file.pdf")


def test_load_documents_from_directory(temp_dir, sample_pdf_path, sample_docx_path):
    pages = load_documents_from_directory(temp_dir)
    # 2 pages from PDF + 1 from DOCX = 3 pages total
    assert len(pages) == 3

    chunks = chunk_documents(pages)
    stats = calculate_document_stats(pages, chunks)

    assert stats["total_documents"] == 2
    assert stats["total_pages"] == 3
    assert stats["total_chunks"] == len(chunks)
    assert "sample_policy.pdf" in stats["chars_per_document"]
    assert "sample_manual.docx" in stats["chars_per_document"]
    assert stats["empty_pages_count"] == 0
